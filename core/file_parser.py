import os
from pathlib import Path

SUPPORTED_EXTS = {".pdf", ".pptx", ".ppt", ".docx", ".doc"}


def extract_text_from_pdf(path):
    import fitz
    doc = fitz.open(path)
    text = []
    for page in doc:
        text.append(page.get_text())
    doc.close()
    return "\n".join(text)


def extract_text_from_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text.append(shape.text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = [cell.text for cell in row.cells]
                    text.append(" | ".join(row_text))
    return "\n".join(text)


def extract_text_from_docx(path):
    from docx import Document
    doc = Document(path)
    text = []
    for para in doc.paragraphs:
        if para.text.strip():
            text.append(para.text)
    return "\n".join(text)


def extract_text(path):
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(path)
    elif ext in (".pptx", ".ppt"):
        return extract_text_from_pptx(path)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(path)
    else:
        raise ValueError(f"\u4e0d\u652f\u6301\u7684\u6587\u4ef6\u683c\u5f0f: {ext}")


def is_supported(path):
    return Path(path).suffix.lower() in SUPPORTED_EXTS
